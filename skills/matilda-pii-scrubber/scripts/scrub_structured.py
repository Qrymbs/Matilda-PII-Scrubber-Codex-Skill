"""Local structured-file adapters preserving nonpersonal values and structure.

Adapters never modify a source file or include raw values in return dictionaries
or error messages. Supported metadata is inspected and preserved unless personal.
"""
from __future__ import annotations

import base64
import csv
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal
from functools import wraps
import json
from pathlib import Path
import re
import warnings
import zipfile

from scrub_core import ScrubContext, ScrubError, classify_field, normalize_field, assess_value, detect_pii_spans, semantic_map_field, requires_context_assessment, is_protected_resource_field

MAX_JSON_BYTES = 512 * 1024 * 1024
MAX_CELL_CHARS = 16 * 1024 * 1024
MAX_OFFICE_UNCOMPRESSED = 1024 * 1024 * 1024
MAX_OFFICE_ENTRIES = 20000


def _guarded(function):
    @wraps(function)
    def wrapped(src, dst, ctx, *args, **kwargs):
        src, dst = Path(src), Path(dst)
        if src.resolve() == dst.resolve():
            raise ScrubError("source_output_collision")
        if dst.exists():
            raise ScrubError("output_already_exists")
        try:
            with warnings.catch_warnings():
                warnings.simplefilter("ignore")
                return function(src, dst, ctx, *args, **kwargs)
        except ScrubError:
            if dst.exists():
                dst.unlink()
            raise
        except Exception:
            if dst.exists():
                dst.unlink()
            raise ScrubError("structured_processing_failed") from None
    return wrapped


def _text_encoding(src: Path) -> str:
    with src.open("rb") as stream:
        prefix = stream.read(4)
    if prefix.startswith((b"\xff\xfe", b"\xfe\xff")):
        return "utf-16"
    if prefix.startswith(b"\xef\xbb\xbf"):
        return "utf-8-sig"
    return "utf-8"


def _safe_names(headers, ctx, location="root"):
    """Inspect schema labels while preserving their exact original values."""
    for index, header in enumerate(headers, 1):
        ctx.scrub(header, field="schema_label", location=f"{location}.c{index}")
    return list(headers)


def _formula_like(value):
    return isinstance(value, str) and bool(re.match(r"^[\s\x00-\x1f]*[=+@]", value))


def _profile(headers, ctx, sheet_name=None):
    raw = {re.sub(r"[^a-z0-9]", "", str(x).lower()) for x in headers if isinstance(x, str)}
    canonical = {normalize_field(x).replace(" ", "") for x in headers}
    profile = "generic_structured"
    if (isinstance(sheet_name, str) and sheet_name.casefold() in {x.casefold() for x in _SAFE_SHEET_NAMES}
            and sheet_name.casefold().startswith(("v", "dv"))):
        profile = "rvtools"
    elif _aws_cost_allocation_schema(headers):
        profile = "aws_cost_allocation"
    elif any(x.startswith("lineitem") for x in raw) and "lineitemunblendedcost" in raw:
        profile = "aws_cur" if any("/" in str(x) for x in headers) else "aws_cur_2"
    elif "costinbillingcurrency" in raw or {"subscriptionid", "resourceid"}.issubset(raw):
        profile = "azure_billing"
    elif "projectid" in raw and ("cost" in raw or "credits" in raw):
        profile = "gcp_billing"
    keys = {"resourceid", "accountid", "usageaccountid", "payeraccountid", "subscriptionid",
            "tenantid", "projectid", "vm", "vmid", "vmname", "instanceid", "hostname",
            "hostid", "databaseid", "diskid", "customerid"}
    mandatory = {
        "aws_cur": {"productcode", "productname", "operation", "lineitemdescription", "lineitemtype", "currencycode", "unit", "usageamount", "unblendedcost", "usagetype"},
        "aws_cur_2": {"productcode", "productname", "operation", "lineitemdescription", "lineitemtype", "currencycode", "unit", "usageamount", "unblendedcost", "usagetype"},
        "azure_billing": {"resourceid", "regionid", "chargeperiodstart", "chargeperiodend", "description", "pricingquantity", "pricingquantitys", "listcost", "listcosts", "skuid"},
        "gcp_billing": {"serviceid", "skuid"},
        "rvtools": {"guestdetaileddata"},
    }.get(profile, set())
    return {"source_profile": profile, "_join_fields_present": sorted(canonical & keys),
            "_mandatory_fields_present": sorted(canonical & mandatory)}


def _final_profile(profile, ctx):
    # Span masks never remove keys or null mandatory values. Keep these legacy
    # fields truthful for report consumers still reading the older names.
    return {"source_profile": profile["source_profile"], "join_keys_removed": False,
            "mandatory_fields_nulled": False, "join_fields_nulled": [],
            "mandatory_fields_nulled_names": []}


# RE2 candidate search is intentionally broader than confirmed core rules.
# Explicit person/prose/tag fields assess all distinct values. Unicode, encoded
# contact markers, attribution labels, profile paths and actor labels are never
# excluded by a technical column name.
ARROW_CANDIDATE = (r"(?i)@|[%&]|\\[ux]|[\{\[]|[^\x00-\x7F]|"
    r"[0-9]{3}-[0-9]{2}-[0-9]{4}|[0-9]{3}[-. ][0-9]{4}|"
    r"full|first|last|person|employee|contact|owner|user|author|creator|refresh|"
    r"prepared|approved|requested|assigned|reviewed|created|phone|telephone|mobile|"
    r"cell|fax|ssn|social|passport|driver|national|home|residential|birth|dob|saved|"
    r"IAMUser|AssumedRole|FederatedUser|/Users/|/home/|\\Users\\")


def _aws_cost_allocation_schema(headers):
    """Recognize the legacy AWS billing schema, never customer banner text."""
    fields = {re.sub(r"[^a-z0-9]", "", str(x).lower()) for x in headers}
    required = {"invoiceid", "payeraccountid", "linkedaccountid", "recordtype",
                "recordid", "productcode", "usagetype", "usagequantity",
                "billingperiodstartdate", "billingperiodenddate"}
    return (required.issubset(fields) and bool(fields & {"totalcost", "unblendedcost", "blendedcost"})
            and len(headers) == len(set(headers)) and all(headers))


@dataclass(frozen=True)
class CsvLayout:
    headers: list
    delimiter: str
    preamble: tuple = ()
    preamble_lines: int = 0

    @property
    def header_row(self):
        return len(self.preamble) + 1

    def field(self, header):
        # RecordID is a billing line key only under this validated schema. A
        # generic personal-record ID elsewhere must not gain this protection.
        if _aws_cost_allocation_schema(self.headers):
            canonical = normalize_field(header).replace(" ", "")
            if canonical == "recordid":
                return "billingrecordid"
            if canonical in {"payeraccountname", "linkedaccountname"}:
                return "cloudaccountname"
            if str(header).casefold() == 'user:name':
                return "resourcename"
        return header


def csv_layout(path, delimiter=None):
    """Locate a normal header or one validated AWS metadata record + header.

    Every prefix record remains content to inspect and compare. No arbitrary
    ragged rows, blank rows, or unknown schema prefixes are skipped.
    """
    csv.field_size_limit(MAX_CELL_CHARS)
    with Path(path).open("r", encoding=_text_encoding(Path(path)), errors="strict", newline="") as source:
        sample = source.read(65536)
        source.seek(0)
        if delimiter is None:
            if Path(path).suffix.lower() == ".tsv":
                delimiter = "\t"
            else:
                try:
                    delimiter = csv.Sniffer().sniff(sample, delimiters=",\t;|").delimiter
                except csv.Error:
                    delimiter = ","
        if delimiter not in {",", "\t", ";", "|"}:
            raise ScrubError("unsupported_delimiter")
        reader = csv.reader(source, delimiter=delimiter, strict=True)
        first = next(reader, None)
        if first is None:
            return CsvLayout([], delimiter)
        first_end_line = reader.line_num
        if not first:
            raise ScrubError("empty_csv_header")
        second = next(reader, None) if len(first) == 1 else None
        if (delimiter == "," and len(first) == 1 and first[0].strip()
                and len(first[0]) <= MAX_CELL_CHARS and second and _aws_cost_allocation_schema(second)):
            return CsvLayout(second, delimiter, (first,), first_end_line)
        return CsvLayout(first, delimiter)


def csv_header(path, delimiter=None):
    layout = csv_layout(path, delimiter)
    return layout.headers, layout.delimiter


def validate_csv_rows(path, layout):
    """Check every record width, including blank trailing records Arrow omits."""
    csv.field_size_limit(MAX_CELL_CHARS)
    rows = 0
    with Path(path).open(encoding=_text_encoding(Path(path)), errors="strict", newline="") as source:
        reader = csv.reader(source, delimiter=layout.delimiter, strict=True)
        for row_index, row in enumerate(reader, 1):
            expected_width = len(layout.preamble[row_index-1]) if row_index < layout.header_row else len(layout.headers)
            if len(row) != expected_width:
                raise ScrubError("ragged_csv_rows")
            if row_index > layout.header_row:
                rows += 1
    return rows


def arrow_csv_reader(path, headers=None, delimiter=None, block_size=8 * 1024 * 1024, preamble_lines=None):
    import pyarrow as pa
    import pyarrow.csv as pacsv
    layout = csv_layout(path, delimiter)
    headers, delimiter = (layout.headers, layout.delimiter) if headers is None else (headers, delimiter or ",")
    if preamble_lines is None:
        preamble_lines = layout.preamble_lines
    return pacsv.open_csv(path, read_options=pacsv.ReadOptions(block_size=block_size, use_threads=False,
            encoding=_text_encoding(Path(path)), skip_rows=preamble_lines), parse_options=pacsv.ParseOptions(delimiter=delimiter, newlines_in_values=True,
                ignore_empty_lines=False),
        convert_options=pacsv.ConvertOptions(column_types={name: pa.string() for name in headers},
            strings_can_be_null=False, quoted_strings_can_be_null=False, null_values=[]))


class BatchCursor:
    """Align different row-group or CSV block boundaries without row sampling."""
    def __init__(self, batches):
        self.batches, self.batch, self.offset = iter(batches), None, 0

    def take(self, count):
        import pyarrow as pa
        chunks = []
        while count:
            if self.batch is None or self.offset == self.batch.num_rows:
                try:
                    self.batch, self.offset = next(self.batches), 0
                except StopIteration:
                    raise ScrubError("output_record_count_mismatch") from None
            size = min(count, self.batch.num_rows - self.offset)
            chunks.append(self.batch.slice(self.offset, size))
            self.offset += size
            count -= size
        return chunks[0] if len(chunks) == 1 else pa.RecordBatch.from_arrays(
            [pa.concat_arrays([chunk.column(i) for chunk in chunks]) for i in range(chunks[0].num_columns)], schema=chunks[0].schema)

    def ensure_end(self):
        if self.batch is not None and self.offset != self.batch.num_rows:
            raise ScrubError("output_record_count_mismatch")
        try:
            next(self.batches)
        except StopIteration:
            return
        raise ScrubError("output_record_count_mismatch")


def map_value_field(key):
    return semantic_map_field(key)


class ArrowPolicy:
    """Bounded-batch masks with exact coordinates and unchanged native types."""
    def __init__(self, ctx, record=True):
        self.ctx, self.record = ctx, record
        self.decision_rows = {}
        for key in ctx.review_decisions:
            match = re.match(r"^r(\d+)\.c(\d+)(?:\.|$)", key)
            if match:
                self.decision_rows.setdefault(int(match.group(2)), set()).add(int(match.group(1)))

    def clear(self):
        # Kept for callers that previously disposed of the assessment cache.
        pass

    def scalar_action(self, value, field, location, protected_by=None):
        # A protected ancestor controls preservation, while the actual child
        # field still provides evidence that the retained content is personal.
        decisions = {location: self.ctx.review_decisions[location]} if location in self.ctx.review_decisions else {}
        context = self.ctx if self.record else ScrubContext(review_decisions=decisions)
        if protected_by:
            found, _ = assess_value(value, field)
            _, reviews = detect_pii_spans(value, field)
            _, protected_reviews = detect_pii_spans(value, protected_by)
            if (found or "type_masking_conflict" in reviews) and "protected_identifier_contains_pii" not in protected_reviews:
                context.note_review("protected_identifier_contains_pii", location)
            elif not found and reviews and "protected_identifier_review" not in protected_reviews:
                context.note_review("protected_identifier_review", location)
        return context.scrub(value, field=protected_by or field, location=location)

    @staticmethod
    def child_field(parent, child, mapping=False):
        if classify_field(parent) == "personal_container" and normalize_field(child).replace(" ", "") == "name":
            return "personname"
        return map_value_field(child) if mapping else child

    def strings(self, array, field, location, rows=None, column=None, path=(), protected_by=None):
        import numpy as np
        import pyarrow as pa
        import pyarrow.compute as pc
        kind = classify_field(field)
        all_values = (requires_context_assessment(field) or kind.startswith("human_") or kind in {"personal_container", "free_text"}
                      or bool(re.search(r"owner|contact|creator|author|tag|label|user.?name", str(field), re.I)))
        candidates = (pc.and_(pc.is_valid(array), pc.not_equal(array, "")) if all_values else
                      pc.fill_null(pc.match_substring_regex(array, ARROW_CANDIDATE), False))
        if rows is not None and column in self.decision_rows:
            candidates = pc.or_(candidates, pa.array(np.isin(rows, list(self.decision_rows[column]))))
        indices = pc.indices_nonzero(candidates).to_numpy(zero_copy_only=False)
        values = pc.take(array, pa.array(indices)).to_pylist()
        replacements, changed_indices = [], []
        if self.record:
            skipped = len(array) - array.null_count - len(indices)
            self.ctx.inspected += skipped
            self.ctx.retained += skipped
        for index, value in zip(indices, values):
            index = int(index)
            if value is None:
                continue
            output = self.scalar_action(value, field, location(index), protected_by)
            if not isinstance(output, str):
                raise ScrubError("string_mask_type_conflict")
            if output != value:
                changed_indices.append(index)
                replacements.append(output)
        if not changed_indices:
            return array
        selected = np.zeros(len(array), dtype=bool)
        selected[changed_indices] = True
        return pc.replace_with_mask(array, pa.array(selected), pa.array(replacements, type=array.type))

    def process(self, array, field, location, rows=None, column=None, path=(), protected_by=None):
        import numpy as np
        import pyarrow as pa
        import pyarrow.compute as pc
        typ = array.type
        if is_protected_resource_field(field):
            protected_by = field
        if pa.types.is_string(typ) or pa.types.is_large_string(typ):
            return self.strings(array, field, location, rows, column, path, protected_by)
        if pa.types.is_dictionary(typ):
            decoded = array.dictionary_decode()
            clean = self.process(decoded, field, location, rows, column, path, protected_by)
            return array if clean.equals(decoded) else pc.cast(pc.dictionary_encode(clean), typ)
        if pa.types.is_struct(typ):
            fields = [self.process(array.field(i), self.child_field(field, child.name),
                      lambda j, i=i: location(j) + f".k{i+1}", rows, column, path+(i,), protected_by)
                      for i, child in enumerate(typ)]
            return pa.StructArray.from_arrays(fields, fields=list(typ), mask=pc.is_null(array))
        if pa.types.is_map(typ):
            offsets = array.offsets.to_numpy(zero_copy_only=False).astype(np.int64)
            start, end = int(offsets[0]), int(offsets[-1])
            offsets = offsets - start
            keys, items = array.keys.slice(start, end-start), array.items.slice(start, end-start)
            parents = np.repeat(np.arange(len(array)), np.diff(offsets))
            ordinals = np.arange(len(items)) - offsets[parents] + 1
            item_loc = lambda index: location(int(parents[index])) + f".k{int(ordinals[index])}"
            # Map keys are schema, so preserve them and record any conflicts.
            # Key and value occupy distinct neutral leaf coordinates.
            key_fields = {}
            for index, key in enumerate(keys.to_pylist()):
                label = key
                if isinstance(key, bytes):
                    try:
                        label = key.decode("utf-8")
                    except UnicodeError:
                        label = None
                        if self.record:
                            self.ctx.warn("retained_binary_parquet_content_uninspected")
                key_fields[key] = label
                if self.record:
                    self.ctx.scrub(label, field="schema_label", location=item_loc(index) + ".k1")
            output = items
            for key in pc.unique(keys).to_pylist():
                selected = pc.equal(keys, key)
                indices = pc.indices_nonzero(selected).to_numpy(zero_copy_only=False)
                clean = self.process(pc.take(items, pa.array(indices)), self.child_field(field, key_fields[key], mapping=True),
                    lambda j: item_loc(int(indices[j])), rows[parents[indices]] if rows is not None else None,
                    column, path+("item",), protected_by)
                try:
                    output = pc.replace_with_mask(output, selected, clean)
                except (pa.ArrowNotImplementedError, pa.ArrowInvalid):
                    expansion = np.maximum(0, np.cumsum(selected.to_numpy(zero_copy_only=False))-1)
                    output = pc.if_else(selected, pc.take(clean, pa.array(expansion)), output)
            return pa.MapArray.from_arrays(pa.array(offsets, type=pa.int32()), keys, output, type=typ, mask=pc.is_null(array))
        if pa.types.is_list(typ) or pa.types.is_large_list(typ) or pa.types.is_fixed_size_list(typ):
            if pa.types.is_fixed_size_list(typ):
                offsets = np.arange(len(array)+1, dtype=np.int64)*typ.list_size
                start = array.offset*typ.list_size
                values = array.values.slice(start, len(array)*typ.list_size)
            else:
                raw = array.offsets.to_numpy(zero_copy_only=False).astype(np.int64)
                values = array.values.slice(int(raw[0]), int(raw[-1]-raw[0]))
                offsets = raw-raw[0]
            parents = np.repeat(np.arange(len(array)), np.diff(offsets))
            ordinals = np.arange(len(values))-offsets[parents]+1
            clean = self.process(values, field, lambda i: location(int(parents[i]))+f".i{int(ordinals[i])}",
                rows[parents] if rows is not None else None, column, path+("item",), protected_by)
            if pa.types.is_fixed_size_list(typ):
                return pa.FixedSizeListArray.from_arrays(clean, type=typ, mask=pc.is_null(array))
            constructor = pa.LargeListArray if pa.types.is_large_list(typ) else pa.ListArray
            return constructor.from_arrays(pa.array(offsets, type=pa.int64() if pa.types.is_large_list(typ) else pa.int32()),
                                           clean, type=typ, mask=pc.is_null(array))
        if pa.types.is_binary(typ) or pa.types.is_large_binary(typ) or pa.types.is_fixed_size_binary(typ):
            for index, value in enumerate(array.to_pylist()):
                if value is None:
                    continue
                has_decision = location(index) in self.ctx.review_decisions
                if self.record and not has_decision:
                    self.ctx.inspected += 1
                    self.ctx.retained += 1
                try:
                    text = value.decode("utf-8")
                except UnicodeError:
                    if self.record:
                        self.ctx.warn("retained_binary_parquet_content_uninspected")
                    continue
                found, reviews = assess_value(text, field)
                if found:
                    reviews = set(reviews) | {"protected_identifier_contains_pii" if protected_by else "type_masking_conflict"}
                if has_decision:
                    # Passing the native value lets the core enforce type-safe
                    # review semantics rather than applying character offsets
                    # to an arbitrary binary encoding.
                    self.scalar_action(value, field, location(index), protected_by)
                elif self.record:
                    for order, reason in enumerate(sorted(reviews)):
                        self.ctx.note_review(reason, location(index), _item=order == 0)
            return array
        if isinstance(typ, pa.BaseExtensionType):
            if self.record:
                self.ctx.warn("retained_extension_parquet_content_uninspected")
            return array
        context_required = classify_field(field).startswith("human_") or requires_context_assessment(field)
        if not context_required and column not in self.decision_rows:
            if self.record:
                self.ctx.inspected += len(array)-array.null_count
                self.ctx.retained += len(array)-array.null_count
            return array
        selected = (np.arange(len(array)) if context_required else
                    np.flatnonzero(np.isin(rows, list(self.decision_rows[column]))))
        if self.record:
            selected_valid = pc.count(pc.take(array, pa.array(selected))).as_py()
            self.ctx.inspected += len(array)-array.null_count-selected_valid
            self.ctx.retained += len(array)-array.null_count-selected_valid
        for index, value in zip(selected, pc.take(array, pa.array(selected)).to_pylist()):
            # Native numeric/date/boolean values are never cast, nulled or
            # reconstructed. Core records confirmed PII as a blocking conflict.
            self.scalar_action(value, field, location(int(index)), protected_by)
        return array


@_guarded
def scrub_csv(src: Path, dst: Path, ctx: ScrubContext, delimiter=None) -> dict:
    """Full streaming scan with exact CSV-record coordinates; no coercion."""
    layout = csv_layout(src, delimiter)
    headers, delimiter = layout.headers, layout.delimiter
    if not headers:
        dst.write_text("",encoding="utf-8")
        return {"rows":0,"input_rows":0,"output_rows":0,"columns":0,"header_rows":0}
    expected_rows = validate_csv_rows(src, layout)
    for row_index, row in enumerate(layout.preamble, 1):
        # The informational prefix is immutable metadata. Human PII or an
        # unresolved identity inside it is preserved and blocks readiness.
        _safe_names(row, ctx, location=f"r{row_index}")
    names=_safe_names(headers,ctx,location=f"r{layout.header_row}")
    fields=[layout.field(header) for header in headers]
    profile=_profile(headers,ctx)
    try:
        import numpy as np
        import pyarrow as pa
        import pyarrow.csv as pacsv
    except ImportError:
        csv.field_size_limit(MAX_CELL_CHARS)
        rows=0
        with src.open("r",encoding=_text_encoding(src),errors="strict",newline="") as source,dst.open("x",encoding="utf-8",newline="") as target:
            reader=csv.reader(source,delimiter=delimiter,strict=True)
            for _ in range(layout.header_row):next(reader)
            writer=csv.writer(target,delimiter=delimiter,lineterminator="\n")
            writer.writerows(layout.preamble);writer.writerow(names)
            for rows,row in enumerate(reader,1):
                if len(row)!=len(headers):raise ScrubError("ragged_csv_rows")
                writer.writerow([ctx.scrub(value,field=field,location=f"r{rows+layout.header_row}.c{i}") for i,(field,value) in enumerate(zip(fields,row),1)])
        return {"rows":rows,"input_rows":rows,"output_rows":rows,"columns":len(headers),"header_rows":1,
                "header_row":layout.header_row,"preamble_rows":len(layout.preamble),
                "coordinate_system":f"CSV header row={layout.header_row}",**_final_profile(profile,ctx)}
    policy=ArrowPolicy(ctx)
    rows,batches=0,0
    schema=pa.schema([pa.field(name,pa.string()) for name in names])
    try:
        with dst.open("xb") as target:
            if layout.preamble:
                import io
                prefix=io.StringIO(newline="")
                csv.writer(prefix,delimiter=delimiter,lineterminator="\n").writerows(layout.preamble)
                target.write(prefix.getvalue().encode("utf-8"))
            with pacsv.CSVWriter(target,schema,write_options=pacsv.WriteOptions(delimiter=delimiter,include_header=True)) as writer:
                for batch in arrow_csv_reader(src,headers,delimiter,preamble_lines=layout.preamble_lines):
                    row_ids=np.arange(rows+layout.header_row+1,rows+batch.num_rows+layout.header_row+1,dtype=np.int64)
                    clean=[policy.process(array,field,lambda j,i=i:f"r{int(row_ids[j])}.c{i}",row_ids,i,(i,))
                           for i,(field,array) in enumerate(zip(fields,batch.columns),1)]
                    writer.write_batch(pa.RecordBatch.from_arrays(clean,schema=schema))
                    rows+=batch.num_rows;batches+=1
                if rows != expected_rows:
                    raise ScrubError("csv_record_count_mismatch")
    finally:
        policy.clear()
    return {"rows":rows,"input_rows":rows,"output_rows":rows,"columns":len(headers),"header_rows":1,
            "header_row":layout.header_row,"preamble_rows":len(layout.preamble),
            "encoding":"utf-8","batches":batches,"coordinate_system":f"CSV header row={layout.header_row}",
            "all_values_scanned":True,"arrow_streaming":True,**_final_profile(profile,ctx)}


def _unique_object(pairs):
    obj = {}
    for key, value in pairs:
        if key in obj:
            raise ScrubError("duplicate_json_keys")
        obj[key] = value
    return obj


def _invalid_constant(value):
    raise ScrubError("nonstandard_json_number")


def _load_json(stream):
    return json.load(stream, object_pairs_hook=_unique_object, parse_constant=_invalid_constant, parse_float=Decimal)


def _write_json(value, stream):
    """Emit native JSON numbers without converting precise decimals to float."""
    if value is None:
        stream.write("null")
    elif value is True:
        stream.write("true")
    elif value is False:
        stream.write("false")
    elif isinstance(value, Decimal):
        if not value.is_finite():
            raise ScrubError("nonstandard_json_number")
        stream.write(str(value))
    elif isinstance(value, int):
        stream.write(str(value))
    elif isinstance(value, float):
        stream.write(json.dumps(value, allow_nan=False))
    elif isinstance(value, str):
        stream.write(json.dumps(value, ensure_ascii=True))
    elif isinstance(value, list):
        stream.write("[")
        for index, item in enumerate(value):
            if index:
                stream.write(",")
            _write_json(item, stream)
        stream.write("]")
    elif isinstance(value, dict):
        stream.write("{")
        for index, (key, item) in enumerate(value.items()):
            if index:
                stream.write(",")
            stream.write(json.dumps(key, ensure_ascii=True))
            stream.write(":")
            _write_json(item, stream)
        stream.write("}")
    else:
        raise ScrubError("unsupported_json_value")


@_guarded
def scrub_json(src: Path, dst: Path, ctx: ScrubContext, jsonl=False) -> dict:
    encoding = _text_encoding(src)
    if not jsonl and src.stat().st_size > MAX_JSON_BYTES:
        raise ScrubError("json_size_limit_exceeded")
    records = 0
    with src.open("r", encoding=encoding, errors="strict") as source, dst.open("x", encoding="utf-8", newline="\n") as output:
        if jsonl:
            for line in source:
                if not line.strip():
                    output.write("\n")
                    continue
                if len(line) > MAX_CELL_CHARS:
                    raise ScrubError("json_record_size_limit_exceeded")
                records += 1
                data = json.loads(line, object_pairs_hook=_unique_object, parse_constant=_invalid_constant, parse_float=Decimal)
                safe = ctx.scrub(data, location=f"r{records}")
                _write_json(safe, output)
                output.write("\n")
        else:
            data = _load_json(source)
            safe = ctx.scrub(data, location="root")
            records = len(data) if isinstance(data, list) else 1
            _write_json(safe, output)
            output.write("\n")
    return {"records": records, "encoding": "utf-8", "native_nulls": True,
            "span_masking": True, "containers_cleared": 0, "leaf_values_cleared": 0,
            "masked_items": ctx.changed}


def _clean_metadata(metadata, ctx, location="root"):
    """Preserve metadata keys and all unmodified bytes; mask UTF-8 text spans."""
    if not metadata:
        return metadata
    clean = {}
    for index, (key, value) in enumerate(metadata.items(), 1):
        try:
            key_text = key.decode("utf-8") if isinstance(key, bytes) else str(key)
            value_text = value.decode("utf-8") if isinstance(value, bytes) else str(value)
        except UnicodeError:
            clean[key] = value
            ctx.warn("retained_binary_parquet_content_uninspected")
            continue
        ctx.scrub(key_text, field="schema_label", location=f"{location}.k{index}.k1")
        sanitized = ctx.scrub(value_text, field=key_text, location=f"{location}.k{index}")
        clean[key] = value if sanitized == value_text else sanitized.encode("utf-8") if isinstance(value, bytes) else sanitized
    return clean


def _clean_arrow_field(column, index, pa, ctx, location="root"):
    ctx.scrub(column.name, field="schema_label", location=f"{location}.c{index}")
    typ = column.type
    if pa.types.is_struct(typ):
        typ = pa.struct([_clean_arrow_field(child, i, pa, ctx, f"{location}.c{index}") for i, child in enumerate(typ, 1)])
    elif pa.types.is_list(typ) or pa.types.is_large_list(typ) or pa.types.is_fixed_size_list(typ):
        child = _clean_arrow_field(typ.value_field, 1, pa, ctx, f"{location}.c{index}")
        typ = pa.large_list(child) if pa.types.is_large_list(typ) else pa.list_(child, typ.list_size) if pa.types.is_fixed_size_list(typ) else pa.list_(child)
    elif pa.types.is_map(typ):
        key = _clean_arrow_field(typ.key_field, 1, pa, ctx, f"{location}.c{index}")
        item = _clean_arrow_field(typ.item_field, 2, pa, ctx, f"{location}.c{index}")
        typ = pa.map_(key, item, keys_sorted=typ.keys_sorted)
    elif isinstance(typ, pa.BaseExtensionType):
        ctx.warn("retained_extension_parquet_content_uninspected")
    return pa.field(column.name, typ, nullable=column.nullable,
                    metadata=_clean_metadata(column.metadata, ctx, f"{location}.c{index}"))


def _retype(array, typ, pa, pc):
    """Rebind changed nested field metadata without converting native values."""
    if array.type.equals(typ, check_metadata=True):
        return array
    if pa.types.is_struct(typ):
        return pa.StructArray.from_arrays([_retype(array.field(i), child.type, pa, pc) for i, child in enumerate(typ)],
                                          fields=list(typ), mask=pc.is_null(array))
    if pa.types.is_map(typ):
        first, last = int(array.offsets[0].as_py()), int(array.offsets[-1].as_py())
        offsets = pc.subtract(array.offsets, pa.scalar(first, type=array.offsets.type))
        return pa.MapArray.from_arrays(offsets, array.keys.slice(first, last-first),
                    _retype(array.items.slice(first, last-first), typ.item_type, pa, pc), type=typ, mask=pc.is_null(array))
    if pa.types.is_list(typ) or pa.types.is_large_list(typ) or pa.types.is_fixed_size_list(typ):
        if pa.types.is_fixed_size_list(typ):
            start = array.offset*typ.list_size
            values = array.values.slice(start, len(array)*typ.list_size)
            return pa.FixedSizeListArray.from_arrays(_retype(values, typ.value_type, pa, pc), type=typ, mask=pc.is_null(array))
        first, last = int(array.offsets[0].as_py()), int(array.offsets[-1].as_py())
        offsets = pc.subtract(array.offsets, pa.scalar(first, type=array.offsets.type))
        constructor = pa.LargeListArray if pa.types.is_large_list(typ) else pa.ListArray
        return constructor.from_arrays(offsets, _retype(array.values.slice(first, last-first), typ.value_type, pa, pc),
                                       type=typ, mask=pc.is_null(array))
    raise ScrubError("parquet_native_type_change_forbidden")


@_guarded
def scrub_parquet(src: Path, dst: Path, ctx: ScrubContext) -> dict:
    try:
        import numpy as np
        import pyarrow as pa
        import pyarrow.compute as pc
        import pyarrow.parquet as pq
    except ImportError:
        raise ScrubError("missing_parquet_dependency") from None
    try:
        parquet = pq.ParquetFile(src)
    except Exception:
        raise ScrubError("parquet_unreadable_or_encrypted") from None
    schema = parquet.schema_arrow
    profile = _profile(schema.names, ctx)
    fields = [_clean_arrow_field(field, i, pa, ctx) for i, field in enumerate(schema, 1)]
    # The embedded Arrow schema and the physical Parquet footer can contain
    # different values or different keys. Inspect both, with distinct locations.
    source_footer = parquet.metadata.metadata or {}
    footer = _clean_metadata({k:v for k,v in source_footer.items() if k != b"ARROW:schema"},
                             ctx, location="footer")
    # Without an embedded schema, schema_arrow.metadata is only a view of the
    # footer. Count and review that physical value once, at its footer location.
    metadata = (_clean_metadata(schema.metadata, ctx)
                if b"ARROW:schema" in source_footer else footer)
    clean_schema = pa.schema(fields, metadata=metadata)
    # Do not let the writer copy schema-only keys into the physical footer or
    # overwrite footer overrides. Regenerate (never copy) the embedded schema.
    if b"ARROW:schema" in source_footer:
        footer[b"ARROW:schema"] = base64.b64encode(clean_schema.serialize().to_pybytes())
    policy = ArrowPolicy(ctx)
    rows, batches = 0, 0
    try:
        with pq.ParquetWriter(dst, clean_schema.remove_metadata(), compression="snappy", store_schema=False) as writer:
            for batch in parquet.iter_batches(batch_size=16384):
                row_ids = np.arange(rows+1, rows+batch.num_rows+1, dtype=np.int64)
                arrays = []
                for column, (field, array, target) in enumerate(zip(schema, batch.columns, fields), 1):
                    clean = policy.process(array, field.name, lambda j,column=column:f"r{int(row_ids[j])}.c{column}", row_ids, column, (column,))
                    arrays.append(_retype(clean, target.type, pa, pc))
                writer.write_table(pa.Table.from_arrays(arrays, schema=clean_schema))
                rows += batch.num_rows
                batches += 1
            if footer:
                writer.add_key_value_metadata(footer)
    finally:
        policy.clear()
        parquet.close()
    incomplete = bool(ctx.warnings & {"retained_binary_parquet_content_uninspected", "retained_extension_parquet_content_uninspected"})
    return {"rows":rows, "input_rows":rows, "output_rows":rows, "columns":len(fields), "batches":batches,
            "incomplete":incomplete, "native_nulls":True, "span_masking":True, "schema_metadata_removed":False,
            "columns_cleared_to_null_type":0, "actual_nullable_fields_changed":0,
            "coordinate_system":"1-based Parquet data records; .kN is map-entry/struct-field ordinal; .iN is list-item ordinal; root.kN is schema metadata; footer.kN is physical footer metadata",
            "all_values_scanned":True, "arrow_streaming":True, **_final_profile(profile,ctx)}


_SAFE_SHEET_NAMES = set("""
vInfo vCPU vMemory vDisk vNetwork vHost vCluster vDatastore vPartition vSnapshot
vTools vCD vFloppy vUSB vHealth vLicense vSource vResourcePool vRP vMultiPath
vHBA vNIC vSwitch vPort dvPort dvport vMetaData vmetadata vSC_VMK vVSwitch vDVSwitch vDPort vDVS vNVRAM
Sheet1 Sheet2 Sheet3 Sheet4 Sheet5 Data data Usage usage Billing billing Costs costs
Summary summary Compute Storage Network Inventory Configuration
""".split())


def _office_manifest(src, ctx):
    if not zipfile.is_zipfile(src):
        raise ScrubError("xlsx_unreadable_or_encrypted")
    limitations=[]
    with zipfile.ZipFile(src) as archive:
        infos=archive.infolist()
        if len(infos)>MAX_OFFICE_ENTRIES or sum(x.file_size for x in infos)>MAX_OFFICE_UNCOMPRESSED:
            raise ScrubError("office_expansion_limit_exceeded")
        if any(x.flag_bits & 1 for x in infos):
            raise ScrubError("encrypted_office_content")
        names=[x.filename.lower() for x in infos]
        if any(any(token in name for token in ("/media/","/embeddings/","vbaproject","/activex/","/ctrlprops/")) for name in names):
            ctx.warn("retained_opaque_workbook_content_uninspected")
        if any(any(token in name for token in ("/embeddings/","vbaproject","/activex/","/ctrlprops/")) for name in names):
            limitations.append("unsupported_embedded_features_may_not_survive_roundtrip")
    return limitations


def _new_workbook():
    try:
        from openpyxl import Workbook
        from openpyxl.packaging.core import DocumentProperties
    except ImportError:
        raise ScrubError("missing_excel_dependency") from None
    workbook=Workbook(write_only=True)
    workbook.properties=DocumentProperties(creator="",lastModifiedBy="",created=datetime(2000,1,1),modified=datetime(2000,1,1))
    return workbook


def _append_values(sheet, values):
    if values and all(value is None for value in values):
        from openpyxl.cell import WriteOnlyCell
        values=[WriteOnlyCell(sheet,value="")]+list(values[1:])
    sheet.append(values)


def _safe_sheet_name(title,index,used,ctx):
    ctx.scrub(title,field="schema_label",location=f"s{index}")
    name=title
    while name in used:
        name+="_"
    used.add(name)
    return name


@_guarded
def scrub_xlsx(src: Path, dst: Path, ctx: ScrubContext) -> dict:
    from scrub_ooxml import scrub_xlsx as faithful_xlsx
    return faithful_xlsx(src,dst,ctx)


@_guarded
def scrub_xls(src: Path, dst: Path, ctx: ScrubContext) -> dict:
    if dst.suffix.lower()!=".xlsx":
        raise ScrubError("xls_requires_xlsx_output")
    try:
        import xlrd
    except ImportError:
        raise ScrubError("missing_legacy_excel_dependency") from None
    try:
        source=xlrd.open_workbook(src,on_demand=True,formatting_info=False)
    except Exception:
        raise ScrubError("xls_unreadable_or_encrypted") from None
    output=_new_workbook()
    used,rows,profiles=set(),0,[]
    count_sheets=source.nsheets
    ctx.warn("legacy_excel_converted_to_xlsx")
    ctx.warn("legacy_excel_formula_inspection_incomplete")
    try:
        for si in range(count_sheets):
            sheet=source.sheet_by_index(si)
            target=output.create_sheet(_safe_sheet_name(sheet.name,si+1,used,ctx))
            if not sheet.nrows:
                continue
            headers=sheet.row_values(0)
            profiles.append(_profile(headers,ctx,sheet.name))
            target.append(_safe_names(headers,ctx,f"s{si+1}.r1"))
            for ri in range(1,sheet.nrows):
                clean=[]
                for ci in range(sheet.ncols):
                    cell=sheet.cell(ri,ci)
                    value=cell.value
                    if cell.ctype==xlrd.XL_CELL_DATE:
                        value=xlrd.xldate.xldate_as_datetime(value,source.datemode)
                    clean.append(ctx.scrub(value,field=headers[ci] if ci<len(headers) else None,location=f"s{si+1}.r{ri+1}.c{ci+1}"))
                _append_values(target,clean)
            rows+=max(0,sheet.nrows-1)
            source.unload_sheet(si)
        output.save(dst)
    finally:
        source.release_resources()
        output.close()
    profiles=[_final_profile(p,ctx) for p in profiles]
    return {"sheets":count_sheets,"rows":rows,"input_rows":rows,"output_rows":rows,"output_extension":".xlsx",
            "reconstructed_values_only":True,"all_legacy_data_cells_cleared":False,
            "formula_inspection_complete":False,"incomplete":True,"format_limitations":["legacy_xls_converted_to_values_only_xlsx_formula_source_and_formatting_not_preserved"],
            "join_keys_removed":any(p["join_keys_removed"] for p in profiles),
            "mandatory_fields_nulled":any(p["mandatory_fields_nulled"] for p in profiles)}
