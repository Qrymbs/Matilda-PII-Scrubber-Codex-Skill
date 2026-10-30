#!/usr/bin/env python3
"""Synthetic processing AND finalization in this interpreter; no customer inputs."""
from __future__ import annotations
import argparse
import base64
import csv
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import zlib

SCRIPTS = Path(__file__).resolve().parent
FORMATS = ('csv', 'xlsx', 'xls', 'parquet')
# Entirely synthetic, fixed expected values; never populated from customer data.
PERSON = 'Avery Stone'
RESOURCE = 'vm-preflight-001'
# A tiny synthetic BIFF8 workbook generated once with test-only xlwt 1.3.0.
# It contains only the same two headers and values as the other fixtures.
# Standard-library decoding adds no XLS preflight dependency for team installs.
XLS_FIXTURE = 'eJztWE1ME0EU/mZpoa2ltFhMwIQ0JKIgJBAvXmCBqJw0qAeNMcGFLtpQWrIWEryIYo8mJp40Xki4eEG9+BM96M2DCUYPJiYmrcaTJxNNPECXN293tSgm9ABoMt9mvnn75r2d130z05l9vRQrzD9oKuI39KIKJTuI6jKdoBL0bqKgdtuWolcHqNgK/xWCAUpktR9Pa1/VyBzKfBeh4b7vBTHwkcpZTCKxxRjgGAwhY+ghFrhDmggaOap65lHmncz32PIZcx9rrjP3kG1BnMGSfqz9oDuKT2st3BaBfO4j9nnPmm404KUcxZdvCMfWj34rZaT/zYZmXxgLoLwNmhnTMtIFxCmBC/huJ4Bv3kx9nlD6rdULkP7HWn3NOvqbmg+YhX2OB3geYXyuki0+HDJyxjIGSAKXHYBlXsxOWaPmcCpJC+3YVDqdMSZMmr3TE52TljmWTp2/kOvs6uom2/5p05pJnMxlM2aIvJ1JHl0zyWt58IeJk6hjOcZTIErL+PLdr2+Ojgzpw6yZ5YXdWf73yIhh44r0IOcIt2g/uZ15P/NVfupulpuY4zRwqW4danCFI3Nsc41bW6mfA4y3+t4yeR/J+S/HHzfnP+ltJC8OFi/FF9/p82ihv6Mk+ctrDh2iQ9y+JfFE92rhLhUfmBv/WDYCWtSN3XaSIuqwghCLMWbnTr4d4dqLdewF22uuvWB7mcperR4P+fX04RdCUFBQUFBQUFBQUNgWOJtZ8GZVbu39cI4H8pwgv+usUClt90cKhU3DCWTpytHB9DAyVFuYqWj87IJfeM8SG/TxvhdKnKLeLYxjhOMYr6hvCTreifLfs2HHaMVd/Q0V91+qJM5N7n8VI5nMZg=='

# Synthetic formula workbook: an opaque formula must remain blocked.
XLS_BLOCKED_FIXTURE = 'eJztWE1oE0EU/maTND+m6aamQiuGULBqbSHFi5d2q6g9KdWLIkLdJlFDY1O2UdCL1ZqjIHhSvBR68VL14g960JsHoWIPgiAkevQkKHhos755u6upVmgObRXm28w3b9+82XnZ92Z2Z9/MxyszDzuq+A0D8KFmh9FUpxNUwt6JDmq3bSl6dYiKrfBfIRyiQDYF8Kz5dVDGUMa7Cg0P/C+JgY9UTmECqXXGfvbBFNKHfmKBu6SJoZ29amXOMG9mvs+Wz5kHWXODuZ9sK+Ik5o0j3XvdLD6hdXJbDPK6j7nPe9b0oQ2vZBZfuSkc2wD2WXmz8G82JP1RzILiNpQbz1lmoYIEBXAW3+wU8NWbqS9SSr++egHSf1+uD66gv6X5gSnYpznBy4hiwSdb/DhglsxFpCFPZdkEWLnJ4gUrkxvJZ6lXpjhZopl78XzvhJU7U8ifPVfqTaf7ImTrTGl92ZRu5lSPEmfRwnKcE16nRXvx3pe3h0eHjRHWTPEy7iz226V/sHFV9qDOMW7RfnI3827ma3zVrSx3MCcoTanuGm5zhUPTbHOdW7tonD2MBWNHnbyT5PLno0+S5U/GLpLnhqqXE3PvjBl00sMnS/3lMY0e0SPu3JZ4ani1cBeGD8ztfywSIU13fbedEIgWLCHCYpzZOZN3R7j2YgV7wfYa39tt7E+coyQff45hEEkhf74BrRWP+I4N4hciUFBQUFBQUFBQUNgIOK+3ziZDvuwH4GwY5M5BftdZolLb6I8UCmuGYyjSUaKN6UGMU23hUkP5swUB4V1LrLKP971Q4jiNbmEMo+zHWENjS9CGT9T/n1V31Bse6m9oePxaI36u8fg/AJ1ox8g='


def run(arguments):
    env = dict(os.environ, PYTHONDONTWRITEBYTECODE='1')
    return subprocess.run([sys.executable, *map(str, arguments)], capture_output=True,
                          text=True, env=env, timeout=120)


def fixture(path, kind):
    row = {'resource_id': RESOURCE, 'fullname': PERSON}
    if kind == 'csv':
        with path.open('w', encoding='utf-8', newline='') as stream:
            writer = csv.writer(stream)
            writer.writerow(row); writer.writerow(row.values())
    elif kind == 'parquet':
        import pyarrow as pa
        import pyarrow.parquet as pq
        pq.write_table(pa.table({key: [value] for key, value in row.items()}), path)
    elif kind == 'xlsx':
        from openpyxl import Workbook
        book = Workbook(); book.active.append(list(row)); book.active.append(list(row.values()))
        book.properties.creator = ''; book.properties.lastModifiedBy = ''
        book.save(path); book.close()
    elif kind == 'xls':
        path.write_bytes(zlib.decompress(base64.b64decode(XLS_FIXTURE)))
    else:
        raise ValueError('unsupported_preflight_format')


def correct_output(path, kind):
    expected = {'resource_id': RESOURCE, 'fullname': 'X' * 11}
    if kind == 'csv':
        with path.open(encoding='utf-8', newline='') as stream:
            return list(csv.reader(stream)) == [list(expected), list(expected.values())]
    if kind == 'parquet':
        import pyarrow.parquet as pq
        return pq.read_table(path).to_pydict() == {key: [value] for key, value in expected.items()}
    if kind == 'xls':
        import xlrd
        with xlrd.open_workbook(path) as book:
            sheet = book.sheet_by_index(0)
            return [sheet.row_values(i) for i in range(sheet.nrows)] == [list(expected), list(expected.values())]
    if kind == 'xlsx':
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True)
        values = list(book.active.values); book.close()
        return values == [tuple(expected), tuple(expected.values())]
    return False


def check_format(root, kind):
    if kind not in FORMATS:
        return {'format': kind, 'passed': False, 'error': 'unsupported_preflight_format'}
    source = root / ('synthetic.' + kind); fixture(source, kind)
    original = source.read_bytes(); out = root / ('run-' + kind)
    process = run([SCRIPTS / 'scrub.py', source, '--output-dir', out])
    if process.returncode != 0:
        return {'format': kind, 'passed': False, 'error': 'synthetic_processing_failed'}
    report = json.loads((out / 'scrub-report.json').read_text())
    record = report['files'][0]
    if not record.get('artifact') or not correct_output(out / record['artifact'], kind) or source.read_bytes() != original:
        return {'format': kind, 'passed': False, 'error': 'independent_expected_output_mismatch'}
    if record.get('output_format') != kind:
        return {'format': kind, 'passed': False, 'error': 'output_format_changed'}
    events = [json.loads(line) for line in (out / record['removal_location_index']).read_text().splitlines()]
    # Location comes from the parser; expected category and spans are independently
    # authored above and compared to the ledger, not accepted from its detector.
    if len(events) != 1 or events[0]['category'] != 'human_names':
        return {'format': kind, 'passed': False, 'error': 'synthetic_mask_count_mismatch'}
    event = events[0]
    if event.get('spans') != [{'start': 0, 'end': 11}]:
        return {'format': kind, 'passed': False, 'error': 'synthetic_mask_span_mismatch'}
    from review_support import COVERAGE_FIELDS
    decisions = {event['location']: {'action': 'mask', 'reason_code': 'explicit_person_name',
                    'spans': [{'start': 0, 'end': 11, 'category': 'human_names'}]}}
    evidence = {'version': 2, 'run_id': report['run_id'], 'files': {
        record['file_id']: {**{key: True for key in COVERAGE_FIELDS}, 'unresolved_items': 0, 'decisions': decisions}}}
    proof = out / 'review-evidence.json'; proof.write_text(json.dumps(evidence), encoding='utf-8')
    final = run([SCRIPTS / 'finalize_report.py', out / 'scrub-report.json', '--review-evidence', proof,
                 '--source', record['file_id'] + '=' + str(source)])
    final_record = json.loads((out / 'scrub-report.json').read_text())['files'][0]
    expected_status = 'ready_for_manual_submission'
    passed = (final.returncode == 0
              and final_record['submission_status'] == expected_status
              and correct_output(out / final_record['artifact'], kind) and source.read_bytes() == original)
    return {'format': kind, 'passed': passed, 'processing_and_finalization_tested': True,
            'expected_submission_status': expected_status,
            **({} if passed else {'error': 'synthetic_finalization_failed'})}


def check_xls(root):
    """Prove a supported native XLS is ready AND an opaque formula is blocked.

    The top-level status describes the blocked fixture, retaining the Version
    1.0 updater's preflight contract. The native case is independently executed
    and required, so older updaters can safely install the new native adapter.
    """
    supported = root / 'xls-supported'; supported.mkdir()
    native = check_format(supported, 'xls')
    source = root / 'opaque-formula.xls'
    source.write_bytes(zlib.decompress(base64.b64decode(XLS_BLOCKED_FIXTURE)))
    original = source.read_bytes(); out = root / 'xls-blocked'
    process = run([SCRIPTS / 'scrub.py', source, '--output-dir', out])
    report = json.loads((out / 'scrub-report.json').read_text())
    record = report['files'][0]
    from review_support import COVERAGE_FIELDS
    evidence = {'version': 2, 'run_id': report['run_id'], 'files': {
        record['file_id']: {**{key: True for key in COVERAGE_FIELDS}, 'unresolved_items': 0, 'decisions': {}}}}
    proof = out / 'review-evidence.json'; proof.write_text(json.dumps(evidence))
    final = run([SCRIPTS / 'finalize_report.py', out / 'scrub-report.json',
                 '--review-evidence', proof, '--source', record['file_id'] + '=' + str(source)])
    final_record = json.loads((out / 'scrub-report.json').read_text())['files'][0]
    blocked_passed = (process.returncode == 2 and final.returncode == 2
                      and final_record['submission_status'] == 'blocked'
                      and 'xls_formula_content_uninspected' in final_record.get('metrics', {}).get('format_limitations', [])
                      and final_record.get('output_format') == 'xls'
                      and bool(final_record.get('artifact'))
                      and (out / final_record['artifact']).read_bytes() == original
                      and source.read_bytes() == original)
    return {'format': 'xls', 'passed': native['passed'] and blocked_passed,
            'processing_and_finalization_tested': True,
            'scenario': 'unsupported_formula_must_remain_blocked',
            'expected_submission_status': 'blocked',
            'native_supported_case': native,
            'unsupported_formula_case': {'passed': blocked_passed, 'expected_submission_status': 'blocked'}}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--formats', nargs='+', choices=FORMATS, default=['csv'])
    args = parser.parse_args()
    os.umask(0o077)
    formats = list(dict.fromkeys(args.formats))
    dependencies = run([SCRIPTS / 'scrub.py', '--check-dependencies', '--formats', *formats])
    if dependencies.returncode:
        print(json.dumps({'passed': False, 'error': 'required_dependencies_unavailable',
                          'formats': formats})); return 2
    results = []
    with tempfile.TemporaryDirectory(prefix='matilda-preflight-') as directory:
        for kind in formats:
            try:
                results.append(check_xls(Path(directory)) if kind == 'xls' else check_format(Path(directory), kind))
            except Exception:
                results.append({'format': kind, 'passed': False, 'error': 'synthetic_preflight_failed'})
    passed = all(row['passed'] for row in results)
    print(json.dumps({'passed': passed, 'synthetic_only': True, 'results': results}, indent=2))
    return 0 if passed else 2


if __name__ == '__main__':
    raise SystemExit(main())
