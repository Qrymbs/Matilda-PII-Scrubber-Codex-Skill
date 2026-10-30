"""Release regressions: synthetic inputs only."""
import csv,json,subprocess,sys,tempfile,unittest
from pathlib import Path
from types import SimpleNamespace
from unittest import mock
import verify_outputs as v
SCRIPTS=Path(__file__).parent
class IndependentVerifier(unittest.TestCase):
 def run_source(self,root,source):
  out=root/'run';proc=subprocess.run([sys.executable,str(SCRIPTS/'scrub.py'),str(source),'--output-dir',str(out)],capture_output=True,text=True)
  self.assertIn(proc.returncode,(0,2),proc.stdout)
  return out,json.loads((out/'scrub-report.json').read_text())
 def choices(self,out,report):
  choices={}
  for line in (out/report['files'][0]['removal_location_index']).read_text().splitlines():
   event=json.loads(line);choice=choices.setdefault(event['location'],{'action':'mask','reason_code':'personal_contact_record','spans':[]})
   choice['spans'].extend(dict(span,category=event['category']) for span in event['spans'])
  return choices
 def test_field_tokens_not_substrings(self):
  for field in ['lineItem/NormalizationFactor','product/normalizationSizeFactor','factor','refactorCount']:self.assertFalse(v._broader_candidate('0.5',field),field)
  for field in ['actor','actor_id','createdBy','userName','employee_id']:self.assertTrue(v._broader_candidate('opaque123',field),field)
  self.assertTrue(v._broader_candidate('casey.rowan@example.invalid','product/ProductName'))
 def test_full_candidate_ledger_beyond_sample_limit(self):
  with tempfile.TemporaryDirectory() as folder:
   root=Path(folder);source=root/'input.csv'
   with source.open('w',newline='') as f:csv.writer(f).writerows([['misc']]+[['Casey Rowan']]*137)
   out,report=self.run_source(root,source);result=v.verify_directory(out,{'file0001':source})['files']['file0001']
   ledger=[json.loads(line) for line in (out/result['independent_candidate_location_index']).read_text().splitlines()]
   self.assertEqual(result['independent_candidates'],137);self.assertEqual(result['unresolved_candidate_count'],137)
   self.assertEqual(len(result['unresolved_candidate_locations']),100);self.assertEqual(len(ledger),137)
   self.assertEqual(ledger[-1]['location'],'r138.c1');self.assertTrue(result['independent_candidate_ledger_complete']);self.assertNotIn('Casey',json.dumps(ledger))
 def test_recorded_spans_are_mechanical_not_semantic_approval(self):
  check=v.ArtifactComparison(SimpleNamespace(max_depth=12),{'output_format':'csv'}, {})
  check.recorded_spans={'r2.c1':[{'start':0,'end':5,'category':'human_names'}]}
  with mock.patch.object(v,'mask_text',side_effect=AssertionError('must not generate expected masks')):check.pair('Casey; cost=1.20','XXXXX; cost=1.20','misc','r2.c1')
  self.assertTrue(check.comparison_ok,check.result);self.assertFalse(check.result['semantic_mask_review_passed']);self.assertEqual(check.result['unreviewed_mask_items'],1);self.assertEqual(check.actions,1)
  bad=v.ArtifactComparison(SimpleNamespace(max_depth=12),{'output_format':'csv'}, {});bad.recorded_spans=check.recorded_spans
  bad.pair('Casey; cost=1.20','XXXXX; cost=1.XX','misc','r2.c1');self.assertFalse(bad.comparison_ok)
 def test_review_decisions_must_match_actual_ledger_spans(self):
  exact={'action':'mask','reason_code':'explicit_person_name','spans':[{'start':0,'end':5,'category':'human_names'}]}
  check=v.ArtifactComparison(SimpleNamespace(max_depth=12),{'output_format':'csv'},{'r2.c1':exact});check.recorded_spans={'r2.c1':exact['spans']}
  check.pair('Casey; cpu=4','XXXXX; cpu=4','misc','r2.c1');self.assertTrue(check.comparison_ok,check.result);self.assertTrue(check.result['semantic_mask_review_passed'])
  wrong=dict(exact,spans=[{'start':0,'end':6,'category':'human_names'}])
  check=v.ArtifactComparison(SimpleNamespace(max_depth=12),{'output_format':'csv'},{'r2.c1':wrong});check.recorded_spans={'r2.c1':exact['spans']}
  check.pair('Casey; cpu=4','XXXXX; cpu=4','misc','r2.c1');self.assertFalse(check.comparison_ok);self.assertFalse(check.result['semantic_mask_review_passed'])
 def test_xlsx_ledger_reconstruction_and_tamper_detection(self):
  import openpyxl
  with tempfile.TemporaryDirectory() as folder:
   root=Path(folder);source=root/'input.xlsx';book=openpyxl.Workbook();book.active.append(['contact_email','cost']);book.active.append(['casey.rowan@example.invalid',12.5]);book.save(source)
   out,report=self.run_source(root,source);choices=self.choices(out,report)
   with mock.patch.object(v.ooxml,'compare_ooxml',side_effect=AssertionError('same-transform forbidden')):result=v.verify_directory(out,{'file0001':source},review_decisions={'file0001':choices})['files']['file0001']
   self.assertTrue(result['source_comparison_passed'],result);self.assertTrue(result['semantic_mask_review_passed'],result)
   artifact=out/report['files'][0]['artifact'];book=openpyxl.load_workbook(artifact);book.active['B2']=12.6;book.save(artifact)
   result=v.verify_directory(out,{'file0001':source},review_decisions={'file0001':choices})['files']['file0001'];self.assertFalse(result['source_comparison_passed']);self.assertIn('ooxml_unexpected_source_output_difference',result['errors'])
 def test_docx_is_denied_without_a_scrubbed_artifact(self):
  from docx import Document
  with tempfile.TemporaryDirectory() as folder:
   root=Path(folder);source=root/'input.docx';book=Document();para=book.add_paragraph();para.add_run('Contact: ');para.add_run('casey.rowan@example.invalid').bold=True;para.add_run('; cpu=4');book.save(source)
   original=source.read_bytes();out,report=self.run_source(root,source);record=report['files'][0]
   self.assertEqual(record['processing_status'],'blocked');self.assertEqual(record['submission_status'],'blocked')
   self.assertIsNone(record['artifact']);self.assertEqual(source.read_bytes(),original)
   self.assertFalse(list(out.rglob('*_SCRUBBED.docx')))
   result=v.verify_directory(out,{'file0001':source})['files']['file0001']
   self.assertFalse(result['source_comparison_passed']);self.assertFalse(result['policy_scan_passed'])
if __name__=='__main__':unittest.main(verbosity=2)
