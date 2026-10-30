#!/usr/bin/env python3
"""Behavioral regression tests: generated fictitious inputs, no customer files."""
from __future__ import annotations
import csv
from decimal import Decimal
import gzip
import importlib.util
import io
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
import zipfile

from scrub_core import ScrubContext, ScrubError
import scrub_structured as structured
import scrub_documents as documents

SCRIPT = Path(__file__).with_name('scrub.py')
PERSON = 'Morgan Taylor'
EMAIL = 'morgan.taylor@example.invalid'

class HumanOnlyBehavior(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='human-pii-synthetic-')
        self.root = Path(self.temp.name)
        self.sources_by_output = {}
    def tearDown(self):
        self.temp.cleanup()
    def csv(self, rows, name='input.csv'):
        path = self.root / name
        with path.open('w', newline='') as stream:
            csv.writer(stream).writerows(rows)
        return path
    def read_csv(self, path):
        with path.open(newline='') as stream:
            return list(csv.reader(stream))
    def launch(self, *paths, name='out', extra=()):
        output = self.root / name
        proc = subprocess.run([sys.executable, str(SCRIPT), *map(str,paths), '--output-dir', str(output), *extra], capture_output=True, text=True)
        report = json.loads((output/'scrub-report.json').read_text())
        top = [r for r in report['files'] if not r.get('parent_file_id')]
        self.sources_by_output[str(output)] = {r['file_id']: str(p) for r,p in zip(top,paths)}
        for content in [proc.stdout, proc.stderr, json.dumps(report), (output/'scrub-report.md').read_text()]:
            self.assertNotIn(EMAIL,content)
            self.assertNotIn(PERSON,content)
        return proc, output, report
    def finalize(self, output):
        report=json.loads((output/'scrub-report.json').read_text())
        evidence={'version':2,'run_id':report['run_id'],'files':{r['file_id']:{'complete_source_review':True,'complete_retained_review':True,'all_changed_values_reviewed':True,'metadata_and_hidden_content_reviewed':True,'independent_candidate_review':True,'unresolved_items':0,'decisions':{}} for r in report['files']}}
        ep=output/'review-evidence.json';ep.write_text(json.dumps(evidence))
        command=[sys.executable,str(SCRIPT.with_name('finalize_report.py')),str(output/'scrub-report.json'),'--review-evidence',str(ep)]
        for fid,path in self.sources_by_output.get(str(output),{}).items():command.extend(['--source',fid+'='+path])
        proc = subprocess.run(command,capture_output=True,text=True)
        return proc,json.loads((output/'scrub-report.json').read_text())
    def test_gcp_billing_schema_and_values_unchanged(self):
        rows=[['Service description','Service ID','SKU description','SKU ID','Usage amount','Usage unit','List cost ($)','Negotiated savings ($)','Savings programs ($)','Other savings ($)','Unrounded subtotal ($)','Subtotal ($)','Percent change in subtotal compared to previous period'],
              ['Compute Engine','6F81-5844-456A','Balanced provisioned capacity','AA11-BB22-CC33','1,234,567.890','gibibyte month','12.345','-1.25','-0.50','0.00','10.595','10.60','12.35%']]
        source=self.csv(rows); before=source.read_bytes(); output=self.root/'clean.csv'; ctx=ScrubContext()
        structured.scrub_csv(source,output,ctx)
        self.assertEqual(self.read_csv(output),rows)
        self.assertEqual(source.read_bytes(),before)
        self.assertEqual(ctx.leaf_values_cleared,0)
    def test_infrastructure_unknown_fields_preserved(self):
        data={'AccountId':'123456789012','ResourceId':'/subscriptions/fake/resourceGroups/prod/providers/compute/vm1','Hostname':'morgan-db-01.corp.invalid','IP':'10.20.30.40','MAC':'00:11:22:33:44:55','Project ID':'business-prod-2026','Company Name':'Example Company','Tenant ID':'11111111-2222-3333-4444-555555555555','Service Account':'billing-agent@project.iam.gserviceaccount.com','Username':'svc_matilda','mystery metric':'1,234,567.89','mystery unit':'gibibyte second','Description':'Database host for the accounting application','Tags':{'environment':'prod','cost center':'1401'},'API key':'machine-key-example','date':'2026-09-01','path':'/var/lib/database'}
        source=self.root/'input.json'; source.write_text(json.dumps(data)); output=self.root/'clean.json'
        structured.scrub_json(source,output,ScrubContext())
        self.assertEqual(json.loads(output.read_text()),data)
    def test_explicit_person_strings_masked_numeric_type_retained_and_blocked(self):
        fields={'Full Name':PERSON,'Email':EMAIL,'Phone':'+1 (202) 555-0147','Date of Birth':'1985-03-17','SSN':'123-45-6789','Employee ID':'EMP-1042','Home Address':'123 Example Street, Sample City','Personal Bank Account':'123456789','Salary':90000}
        source=self.root/'input.json'; source.write_text(json.dumps(fields)); output=self.root/'clean.json'
        ctx=ScrubContext(); structured.scrub_json(source,output,ctx)
        result=json.loads(output.read_text())
        self.assertEqual(set(result),set(fields))
        for key,value in fields.items():
            self.assertEqual(result[key], 'X'*len(value) if isinstance(value,str) else value)
            self.assertIs(type(result[key]),type(value))
        self.assertEqual(ctx.changed,len(fields)-1)
        self.assertGreater(ctx.review_counts['type_masking_conflict'],0)
    def test_name_field_does_not_erase_machine_name(self):
        source=self.csv([['Name','VM name','Account ID','SKU ID'],['morgan-db-01','morgan-db-01','123456789012','AA11-BB22-CC33']])
        output=self.root/'clean.csv'; structured.scrub_csv(source,output,ScrubContext())
        self.assertEqual(self.read_csv(output),self.read_csv(source))
    def test_company_owners_and_phone_shaped_catalog_ids_are_not_people(self):
        data={'Owner':'Example Company','Contact':'Operations Team','Service ID':'123-45-6789','SKU ID':'202-555-0147','ResourceId':'vm-prod-01','Owner Name':'Example Company LLC'}
        self.assertEqual(ScrubContext().scrub(data),data)
    def test_person_owner_and_credentials_masked_machine_values_exact(self):
        data={'Owner':PERSON,'Personal Password':'synthetic-person-password','machine secret':'synthetic-machine-secret','service username':'svc_matilda'}
        result=ScrubContext().scrub(data)
        self.assertEqual(result['Owner'],'X'*len(PERSON)); self.assertEqual(result['Personal Password'],'X'*len(data['Personal Password']))
        self.assertEqual(result['machine secret'],data['machine secret']); self.assertEqual(result['service username'],'svc_matilda')
    def test_description_masks_only_person_span(self):
        source=self.csv([['Description','Notes','ResourceId'],['Production DB; contact '+EMAIL,'Regular monthly cost report','/projects/prod/resources/db1']])
        output=self.root/'clean.csv'; structured.scrub_csv(source,output,ScrubContext())
        self.assertEqual(self.read_csv(output)[1],['Production DB; contact '+'X'*len(EMAIL),'Regular monthly cost report','/projects/prod/resources/db1'])
    def test_labeled_person_in_free_text_preserves_label_and_other_lines(self):
        source=self.root/'input.txt'; source.write_text('Database host is healthy.\nFull name: '+PERSON+'\nContact email: '+EMAIL+'\nDisk size: 500 GiB\n')
        output=self.root/'clean.txt'; documents.scrub_text(source,output,ScrubContext())
        lines=output.read_text().splitlines()
        self.assertEqual(lines,['Database host is healthy.','Full name: '+'X'*len(PERSON),'Contact email: '+'X'*len(EMAIL),'Disk size: 500 GiB'])
    def test_human_email_sentence_punctuation_and_attribution(self):
        for text in ['Assigned to '+PERSON+'; email '+EMAIL+'.', 'Reviewed by '+PERSON+' during onboarding.', 'Contact ('+EMAIL+').', 'Contact '+EMAIL+', then continue.']:
            result=ScrubContext().scrub(text,field='Notes')
            self.assertEqual(result,text.replace(PERSON,'X'*len(PERSON)).replace(EMAIL,'X'*len(EMAIL)))
            self.assertEqual(len(result),len(text))
    def test_mixed_pii_report_lists_each_type_without_raw_values(self):
        ctx=ScrubContext()
        phone='+1 (202) 555-0147'
        text='Full name: '+PERSON+'; email: '+EMAIL+'; Phone: '+phone
        self.assertEqual(ctx.scrub(text,field='Notes'),text.replace(PERSON,'X'*len(PERSON)).replace(EMAIL,'X'*len(EMAIL)).replace(phone,'X'*len(phone)))
        self.assertTrue({'human_names','human_email_addresses','human_phone_numbers'}.issubset(ctx.counts))
        self.assertEqual(ctx.leaf_values_cleared,0)
        self.assertEqual(ctx.changed,1)
        self.assertEqual(ctx.masked_spans,3)
        self.assertEqual(ctx.masked_characters,len(PERSON)+len(EMAIL)+len(phone))
        self.assertNotIn(PERSON,json.dumps(ctx.safe_report())); self.assertNotIn(EMAIL,json.dumps(ctx.safe_report()))
    def test_safe_multiline_description_preserved(self):
        value='Compute cluster report\nRegion: us-east-1\nUsage: 1,000 gibibyte hours'
        source=self.csv([['Description','cost'],[value,'1,000.00']]); output=self.root/'clean.csv'
        structured.scrub_csv(source,output,ScrubContext())
        self.assertEqual(self.read_csv(output),self.read_csv(source))
    def test_nested_pii_does_not_clear_infrastructure_parent(self):
        data={'ResourceId':'vm-prod-01','properties':{'contact_email':EMAIL,'disk_gib':120,'note':'Routine maintenance'},'tags':{'environment':'prod','owner_email':EMAIL}}
        source=self.root/'input.json'; source.write_text(json.dumps(data)); output=self.root/'clean.json'
        structured.scrub_json(source,output,ScrubContext()); result=json.loads(output.read_text())
        self.assertEqual(result['ResourceId'],data['ResourceId'])
        self.assertEqual(result['properties']['disk_gib'],120)
        self.assertEqual(result['properties']['note'],'Routine maintenance')
        self.assertEqual(result['properties']['contact_email'],'X'*len(EMAIL))
        self.assertEqual(result['tags']['environment'],'prod')
        self.assertEqual(result['tags']['owner_email'],'X'*len(EMAIL))
    def test_json_exact_decimal_and_large_cloud_ids(self):
        source=self.root/'input.json'; source.write_text('{"Usage amount":12345678901234567890,"cost":12.123456789012345678901234,"Full Name":"'+PERSON+'"}')
        output=self.root/'clean.json'; structured.scrub_json(source,output,ScrubContext())
        result=json.loads(output.read_text(),parse_float=Decimal)
        self.assertEqual(result['Usage amount'],12345678901234567890)
        self.assertEqual(result['cost'],Decimal('12.123456789012345678901234'))
        self.assertEqual(result['Full Name'],'X'*len(PERSON))
    def test_jsonl_and_utf16_tsv(self):
        source=self.root/'input.jsonl'; source.write_text(json.dumps({'Email':EMAIL,'ResourceId':'vm-1'})+'\n'+json.dumps({'Email':EMAIL,'ResourceId':'vm-2'})+'\n')
        output=self.root/'clean.jsonl'; structured.scrub_json(source,output,ScrubContext(),jsonl=True)
        self.assertEqual([json.loads(x) for x in output.read_text().splitlines()],[{'Email':'X'*len(EMAIL),'ResourceId':'vm-1'},{'Email':'X'*len(EMAIL),'ResourceId':'vm-2'}])
        source=self.root/'input.tsv'; source.write_text('Email\tunknown_unit\n'+EMAIL+'\tgibibyte second\n',encoding='utf-16')
        output=self.root/'clean.tsv'; structured.scrub_csv(source,output,ScrubContext(),delimiter='\t')
        with output.open(newline='') as stream:self.assertEqual(list(csv.reader(stream,delimiter='\t'))[1],['X'*len(EMAIL),'gibibyte second'])
    def test_duplicate_json_and_ragged_csv_blocked(self):
        source=self.root/'duplicate.json'; source.write_text('{"cost":1,"cost":2}'); output=self.root/'clean.json'
        with self.assertRaises(ScrubError): structured.scrub_json(source,output,ScrubContext())
        self.assertFalse(output.exists())
        source=self.root/'ragged.csv'; source.write_text('cost,Email\n1,contact,extra\n'); output=self.root/'clean.csv'
        with self.assertRaises(ScrubError): structured.scrub_csv(source,output,ScrubContext())
        self.assertFalse(output.exists())
    @unittest.skipUnless(importlib.util.find_spec('pyarrow'),'pyarrow unavailable')
    def test_parquet_nested_values_and_types(self):
        import pyarrow as pa
        import pyarrow.parquet as pq
        data=pa.table({'ResourceId':['vm1'],'IP':['10.0.0.1'],'Email':[EMAIL],'cost':pa.array([Decimal('12.123456')],type=pa.decimal128(18,6)),'properties':pa.array([{'disk':32,'contact_email':EMAIL}])})
        source=self.root/'input.parquet'; pq.write_table(data,source); output=self.root/'clean.parquet'
        structured.scrub_parquet(source,output,ScrubContext()); result=pq.read_table(output).to_pylist()[0]
        self.assertEqual(result['ResourceId'],'vm1'); self.assertEqual(result['IP'],'10.0.0.1'); self.assertEqual(result['Email'],'X'*len(EMAIL))
        self.assertEqual(result['cost'],Decimal('12.123456')); self.assertEqual(result['properties']['disk'],32); self.assertEqual(result['properties']['contact_email'],'X'*len(EMAIL))
    @unittest.skipUnless(importlib.util.find_spec('openpyxl'),'openpyxl unavailable')
    def test_xlsx_safe_features_and_hidden_human_pii(self):
        from openpyxl import Workbook,load_workbook
        from openpyxl.comments import Comment
        source=self.root/'input.xlsx'; workbook=Workbook(); sheet=workbook.active; sheet.title='Custom billing'
        sheet.append(['ResourceId','Description','cost','Full Name']); sheet.append(['vm-prod-01','Routine billing',12.5,PERSON]); sheet.append(['vm-prod-02','Contact '+EMAIL,7.5,None])
        sheet['C4']='=SUM(C2:C3)'; sheet['B2'].comment=Comment('Database maintenance','Automation'); sheet['A2'].hyperlink='https://console.example.invalid/vm-prod-01'
        hidden=workbook.create_sheet('Hidden infrastructure'); hidden.sheet_state='hidden'; hidden.append(['Email','IP']); hidden.append([EMAIL,'10.0.0.1']); workbook.save(source)
        before=source.read_bytes(); output=self.root/'clean.xlsx'; structured.scrub_xlsx(source,output,ScrubContext()); result=load_workbook(output,data_only=False)
        self.assertEqual(result.sheetnames,workbook.sheetnames); self.assertEqual(result.worksheets[0]['A2'].value,'vm-prod-01'); self.assertEqual(result.worksheets[0]['D2'].value,'X'*len(PERSON)); self.assertEqual(result.worksheets[0]['B3'].value,'Contact '+'X'*len(EMAIL))
        self.assertEqual(result.worksheets[0]['C4'].value,'=SUM(C2:C3)'); self.assertEqual(result.worksheets[0]['B2'].comment.text,'Database maintenance'); self.assertEqual(result.worksheets[1]['B2'].value,'10.0.0.1'); self.assertEqual(result.worksheets[1]['A2'].value,'X'*len(EMAIL)); self.assertEqual(source.read_bytes(),before)
    @unittest.skipUnless(importlib.util.find_spec('xlrd') and importlib.util.find_spec('xlwt') and importlib.util.find_spec('openpyxl'),'legacy Excel libraries unavailable')
    def test_xls_preserves_values_and_discloses_inspection_limit(self):
        import xlwt
        from openpyxl import load_workbook
        source=self.root/'input.xls'; workbook=xlwt.Workbook(); sheet=workbook.add_sheet('Infrastructure')
        for col,value in enumerate(['ResourceId','Email','cost']): sheet.write(0,col,value)
        for col,value in enumerate(['vm-prod-01',EMAIL,12.5]): sheet.write(1,col,value)
        workbook.save(str(source)); output=self.root/'clean.xlsx'; metrics=structured.scrub_xls(source,output,ScrubContext()); result=load_workbook(output).active
        self.assertEqual(result['A2'].value,'vm-prod-01'); self.assertEqual(result['B2'].value,'X'*len(EMAIL)); self.assertEqual(result['C2'].value,12.5)
        self.assertTrue(metrics.get('incomplete'))
        _,batch,report=self.launch(source)
        self.assertEqual(report['files'][0]['privacy_status'],'partial')
        _,final=self.finalize(batch)
        self.assertEqual(final['files'][0]['submission_status'],'blocked')
    @unittest.skipUnless(importlib.util.find_spec('openpyxl'),'openpyxl unavailable')
    def test_required_formula_cache_preserved(self):
        from openpyxl import Workbook
        from xml.etree import ElementTree as ET
        source=self.root/'input.xlsx'; workbook=Workbook(); workbook.active.append(['SubscriptionId','ResourceId','ListCosts']); workbook.active.append(['subscription-prod','vm-prod-01','=1+1']); workbook.save(source); workbook.close()
        with zipfile.ZipFile(source) as package: parts={name:package.read(name) for name in package.namelist()}
        namespace='{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
        tree=ET.fromstring(parts['xl/worksheets/sheet1.xml'])
        cell=next(item for item in tree.iter(namespace+'c') if item.get('r')=='C2')
        cached=cell.find(namespace+'v')
        if cached is None: cached=ET.SubElement(cell,namespace+'v')
        cached.text='2'; parts['xl/worksheets/sheet1.xml']=ET.tostring(tree)
        with zipfile.ZipFile(source,'w',compression=zipfile.ZIP_DEFLATED) as package:
            for name,raw in parts.items(): package.writestr(name,raw)
        _,batch,report=self.launch(source)
        self.assertEqual(report['files'][0]['metrics'].get('known_required_formula_cache_loss',0),0)
        self.assertNotEqual(report['files'][0]['matilda_compatibility'],'not_ready')
        with zipfile.ZipFile(batch/report['files'][0]['artifact']) as package:
            saved=ET.fromstring(package.read('xl/worksheets/sheet1.xml'))
        saved_cell=next(item for item in saved.iter(namespace+'c') if item.get('r')=='C2')
        self.assertEqual(saved_cell.find(namespace+'v').text,'2')
        self.assertEqual(saved_cell.find(namespace+'f').text,'1+1')
    @unittest.skipUnless(importlib.util.find_spec('openpyxl') and importlib.util.find_spec('PIL'),'image spreadsheet libraries unavailable')
    def test_retained_uninspected_workbook_image_blocks_release(self):
        from PIL import Image
        from openpyxl import Workbook
        from openpyxl.drawing.image import Image as ExcelImage
        picture=io.BytesIO()
        with Image.new('RGB',(2,2),'white') as image: image.save(picture,format='PNG')
        picture.seek(0)
        source=self.root/'input.xlsx'; workbook=Workbook(); workbook.active.append(['ResourceId']); workbook.active.append(['vm-prod-01']); workbook.active.add_image(ExcelImage(picture),'C1'); workbook.save(source); workbook.close()
        _,batch,report=self.launch(source)
        self.assertEqual(report['files'][0]['privacy_status'],'partial')
        _,final=self.finalize(batch)
        self.assertEqual(final['files'][0]['submission_status'],'blocked')
    @unittest.skipUnless(importlib.util.find_spec('docx'),'python-docx unavailable')
    def test_docx_masked_paragraph_and_table_cell_keep_safe_fragments(self):
        from docx import Document
        source=self.root/'input.docx'; document=Document(); document.add_paragraph('Production cluster runs Linux.'); document.add_paragraph('Production cluster contact: '+EMAIL); table=document.add_table(rows=2,cols=2); table.cell(0,0).text='Description'; table.cell(0,1).text='ResourceId'; table.cell(1,0).text='Maintenance by '+EMAIL+'\nPreserve this safe fragment of the cell'; table.cell(1,1).text='vm-prod-01'; document.save(source)
        output=self.root/'clean.docx'; documents.scrub_docx(source,output,ScrubContext()); result=Document(output)
        self.assertEqual(result.paragraphs[0].text,'Production cluster runs Linux.'); self.assertEqual(result.paragraphs[1].text,'Production cluster contact: '+'X'*len(EMAIL)); self.assertEqual(result.tables[0].cell(1,0).text,'Maintenance by '+'X'*len(EMAIL)+'\nPreserve this safe fragment of the cell'); self.assertEqual(result.tables[0].cell(1,1).text,'vm-prod-01')
    @unittest.skipUnless(importlib.util.find_spec('pypdf') and importlib.util.find_spec('reportlab'),'PDF libraries unavailable')
    def test_pdf_retains_all_bytes_and_blocks_unsupported_masking(self):
        from reportlab.pdfgen import canvas
        from pypdf import PdfReader
        source=self.root/'input.pdf'; painter=canvas.Canvas(str(source)); painter.drawString(50,750,'Resource vm-prod-01: 1,234 gibibyte hours'); painter.drawString(50,730,'Contact: '+EMAIL); painter.save()
        output=self.root/'clean.pdf'; ctx=ScrubContext(); metrics=documents.scrub_pdf(source,output,ctx); text='\n'.join(page.extract_text() for page in PdfReader(output).pages)
        self.assertEqual(output.read_bytes(),source.read_bytes())
        self.assertIn('vm-prod-01',text); self.assertIn('1,234 gibibyte hours',text); self.assertIn(EMAIL,text)
        self.assertTrue(metrics['incomplete']); self.assertTrue(metrics['layout_preserved'])
        self.assertEqual(ctx.changed,0); self.assertGreater(ctx.review_counts['unsupported_content_uninspected'],0)
        _,batch,_=self.launch(source)
        _,final=self.finalize(batch)
        self.assertEqual(final['files'][0]['submission_status'],'blocked')
    def test_reports_explain_human_types_and_reasons(self):
        source=self.csv([['Full Name','Email','ResourceId'],[PERSON,EMAIL,'vm-prod-01']]); _,output,report=self.launch(source)
        self.assertNotIn('customer_infrastructure_identifiers',report['privacy_scope'])
        self.assertNotIn('secrets',report['privacy_scope'])
        summary=(output/'scrub-report.md').read_text().lower()
        self.assertIn('why',summary); self.assertIn('human',summary)
        self.assertEqual(report['files'][0]['processing_status'],'completed')
        self.assertEqual(report['files'][0]['privacy_status'],'requires_context_review')
    def test_infrastructure_retention_passes_finalization(self):
        source=self.csv([['ResourceId','AccountId','Description','cost','Usage unit'],['vm-prod-01','123456789012','Production database','1,234.50','gibibyte month']]); _,output,report=self.launch(source); proc,final=self.finalize(output)
        self.assertEqual(proc.returncode,0,proc.stdout+proc.stderr)
        self.assertEqual(final['files'][0]['submission_status'],'ready_for_manual_submission')
        self.assertNotEqual(final['files'][0]['matilda_compatibility'],'not_ready')
    def test_verifier_rejects_reintroduced_human_pii(self):
        source=self.csv([['Email','ResourceId'],[EMAIL,'vm-prod-01']]); _,output,report=self.launch(source); artifact=output/report['files'][0]['artifact']
        rows=self.read_csv(artifact); rows[1][0]=EMAIL
        with artifact.open('w',newline='') as stream: csv.writer(stream).writerows(rows)
        _,final=self.finalize(output)
        self.assertEqual(final['files'][0]['submission_status'],'blocked')
    def test_gzip_zip_require_extraction_without_creating_scrubbed_copies(self):
        source=self.csv([['Full Name','ResourceId','cost'],[PERSON,'vm-prod-01','12.50']]); archive=self.root/'input.zip'
        with zipfile.ZipFile(archive,'w') as package: package.writestr('customer-data.csv.gz',gzip.compress(source.read_bytes(),mtime=0))
        archive_bytes=archive.read_bytes();proc,output,report=self.launch(archive); parent=report['files'][0]
        self.assertEqual(proc.returncode,2);self.assertIsNone(parent['artifact'])
        self.assertEqual(archive.read_bytes(),archive_bytes)
        self.assertIn('extract',(output/'scrub-report.md').read_text().lower())
        self.assertFalse(list(output.rglob('*_SCRUBBED*')))
        self.assertEqual(parent['processing_status'],'blocked')
        self.assertEqual(parent['metrics'].get('masked_items',0),0)
        compressed=self.root/'input.csv.gz'; compressed.write_bytes(gzip.compress(source.read_bytes(),mtime=0))
        compressed_bytes=compressed.read_bytes();_,gzip_output,gzip_report=self.launch(compressed,name='gzip-out')
        gzip_record=gzip_report['files'][0]
        self.assertIsNone(gzip_record['artifact']);self.assertEqual(compressed.read_bytes(),compressed_bytes)
        self.assertIn('extract',(gzip_output/'scrub-report.md').read_text().lower())
        self.assertFalse(list(gzip_output.rglob('*_SCRUBBED*')))
        self.assertEqual(gzip_record['submission_status'],'blocked')
        proc,final=self.finalize(output)
        self.assertEqual(proc.returncode,2,proc.stdout+proc.stderr)
        self.assertTrue(all(r['submission_status']=='blocked' for r in final['files']))
        self.assertFalse(final['files'][0]['verification']['source_comparison_passed'])
    def test_batch_unsupported_archive_traversal_and_source_guards(self):
        source=self.csv([['Full Name','cost'],[PERSON,'12.5']]); before=source.read_bytes(); bad=self.root/'input.binary'; bad.write_bytes(b'not a supported format')
        proc,_,report=self.launch(source,bad); self.assertEqual(proc.returncode,2); self.assertEqual(report['files'][0]['processing_status'],'completed'); self.assertEqual(report['files'][0]['privacy_status'],'requires_context_review'); self.assertEqual(report['files'][1]['privacy_status'],'blocked'); self.assertEqual(source.read_bytes(),before)
        archive=self.root/'unsafe.zip'
        with zipfile.ZipFile(archive,'w') as package: package.writestr('../outside.txt','no execution')
        _,_,report=self.launch(archive,name='archive-out'); self.assertEqual(report['files'][0]['submission_status'],'blocked'); self.assertFalse((self.root/'outside.txt').exists())
        alias=self.root/'alias.csv'; alias.symlink_to(source); _,_,report=self.launch(alias,name='alias-out'); self.assertEqual(report['files'][0]['privacy_status'],'blocked')
        proc=subprocess.run([sys.executable,str(SCRIPT),str(source),'--output-dir',str(self.root)],capture_output=True,text=True); self.assertEqual(proc.returncode,2)

if __name__=='__main__':
    unittest.main(verbosity=2)
